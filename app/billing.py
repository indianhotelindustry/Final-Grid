"""
Billing Blueprint  (/billing/*)
================================
Covers:
  - GST summary report
  - Payment void request creation (FrontDesk)
  - Void request approval/rejection dashboard (Manager/Admin)
  - Direct void (Manager/Admin single-step)
  - Pending void request count for navbar badge (via context processor in __init__)
"""
from functools import wraps

from flask import Blueprint, flash, jsonify, redirect, render_template, request, url_for
from flask_login import current_user, login_required

from app.auth import role_required
from app.models import Payment, VoidRequest, db
from app import gst_service, payment_void_service, limiter

billing_bp = Blueprint('billing', __name__, url_prefix='/billing')


# ---------------------------------------------------------------------------
# Role helpers
# ---------------------------------------------------------------------------

def _manager_required(f):
    @wraps(f)
    @login_required
    def decorated(*args, **kwargs):
        if current_user.role not in ('Admin', 'Manager'):
            flash('Manager or Admin access required.', 'danger')
            return redirect(url_for('main.dashboard'))
        return f(*args, **kwargs)
    return decorated


# ---------------------------------------------------------------------------
# GST Report
# ---------------------------------------------------------------------------

@billing_bp.route('/gst-report')
@role_required('Admin', 'Manager', 'Accountant')
def gst_report():
    from datetime import date
    from_str = request.args.get('from_date', '')
    to_str   = request.args.get('to_date', '')

    try:
        from_date = date.fromisoformat(from_str) if from_str else date.today().replace(day=1)
        to_date   = date.fromisoformat(to_str)   if to_str   else date.today()
    except ValueError:
        from_date = date.today().replace(day=1)
        to_date   = date.today()

    data = gst_service.get_gst_report(from_date, to_date)
    hotel_gstin   = gst_service.get_hotel_gstin()
    hotel_state   = gst_service.get_hotel_state_code()

    return render_template(
        'billing/gst_report.html',
        data=data,
        hotel_gstin=hotel_gstin,
        hotel_state=hotel_state,
    )


# ---------------------------------------------------------------------------
# Invoice Register
# ---------------------------------------------------------------------------

@billing_bp.route('/invoice-register')
@role_required('Admin', 'Manager', 'FrontDesk', 'Accountant')
def invoice_register():
    """
    Central invoice register — all checked-out reservation bills.
    Source of truth: Reservation.status == 'CheckedOut'.
    Billing amounts derived from payments + extra_charges (eager-loaded).
    Checkout user from AuditLog batch query.
    """
    from datetime import date, timedelta
    from decimal import Decimal
    from flask import request, render_template, send_file, Response
    from sqlalchemy import func
    from app.models import (Reservation, Room, Guest, Payment, ExtraCharge,
                            AuditLog, User, PaymentMode, db)
    import io

    today    = date.today()
    quick    = request.args.get('tab',    'today').strip()   # today/unsettled/partial/settled/all
    search   = request.args.get('search', '').strip()
    room_f   = request.args.get('room',   '').strip()
    btype_f  = request.args.get('btype',  '').strip()
    fmt      = request.args.get('format', 'html')

    # Date range defaults: 'today' tab = today only; others = last 30 days
    _default_from = today.isoformat() if quick == 'today' else (today - timedelta(days=29)).isoformat()
    from_str = request.args.get('from', _default_from)
    to_str   = request.args.get('to',   today.isoformat())
    try:
        from_date = date.fromisoformat(from_str)
        to_date   = date.fromisoformat(to_str)
    except ValueError:
        from_date = to_date = today

    # ── Base query: CheckedOut + eager-load billing relationships ─────────────
    q = (Reservation.query
         .options(db.joinedload(Reservation.payments).joinedload(Payment.payment_mode),
                  db.joinedload(Reservation.extra_charges),
                  db.joinedload(Reservation.guest),
                  db.joinedload(Reservation.room))
         .filter(Reservation.status == 'CheckedOut'))

    # Date filter on checked_out_at
    if quick == 'today':
        q = q.filter(func.date(Reservation.checked_out_at) == today)
    else:
        q = q.filter(func.date(Reservation.checked_out_at) >= from_date,
                     func.date(Reservation.checked_out_at) <= to_date)

    # Optional filters
    if search:
        _like = f'%{search}%'
        q = (q.join(Reservation.guest, isouter=True)
              .filter(db.or_(Guest.name.ilike(_like), Guest.phone.ilike(_like))))
    if room_f:
        q = (q.join(Reservation.room, isouter=True)
              .filter(Room.room_number.ilike(f'%{room_f}%')))
    if btype_f:
        q = q.filter(Reservation.booking_type == btype_f)

    all_reservations = q.order_by(Reservation.checked_out_at.desc()).all()

    # ── Billing helpers (in-memory — relationships already loaded) ────────────
    def _billing(res):
        nights  = (res.departure_date - res.arrival_date).days
        room_ch = Decimal(str(res.rate_per_night or 0)) * Decimal(nights)
        # Exclude night-audit room_rent rows — room_ch already covers the
        # room revenue. Without this filter the invoice register would
        # show every checked-out reservation at ~2× the actual total.
        extra   = sum(
            Decimal(str(ec.amount or 0))
            for ec in res.extra_charges
            if (ec.charge_type or '') != 'room_rent'
        )
        discount = Decimal(str(getattr(res, 'discount_amount', None) or 0))
        total   = room_ch + extra - discount
        paid    = sum(Decimal(str(p.amount or 0)) for p in res.payments if not p.is_voided)
        balance = total - paid
        return {'total': float(total), 'paid': float(paid), 'balance': float(balance)}

    def _status(b):
        from app.financial import SETTLEMENT_TOLERANCE
        _tol = float(SETTLEMENT_TOLERANCE)
        if b['balance'] <= _tol:    return 'Settled'
        if b['paid']    <= _tol:    return 'Unsettled'
        return 'Partial'

    def _modes(res):
        m = {}
        for p in res.payments:
            if not p.is_voided and p.payment_mode:
                m[p.payment_mode.name] = m.get(p.payment_mode.name, 0.0) + float(p.amount)
        return m

    # ── Build rows (all, before status filter) ────────────────────────────────
    all_rows = []
    for res in all_reservations:
        b = _billing(res)
        all_rows.append({
            'reservation': res,
            'billing':     b,
            'status':      _status(b),
            'modes':       _modes(res),
        })

    # ── Summary cards (from date-filtered all_rows, before status filter) ─────
    summary = {
        'count':          len(all_rows),
        'total_billed':   round(sum(r['billing']['total']   for r in all_rows), 2),
        'total_collected':round(sum(r['billing']['paid']    for r in all_rows), 2),
        'total_pending':  round(sum(r['billing']['balance'] for r in all_rows), 2),
        'unsettled_count':sum(1 for r in all_rows if r['status'] == 'Unsettled'),
        'partial_count':  sum(1 for r in all_rows if r['status'] == 'Partial'),
        'settled_count':  sum(1 for r in all_rows if r['status'] == 'Settled'),
    }

    # ── Apply status tab filter ───────────────────────────────────────────────
    if   quick == 'unsettled': rows = [r for r in all_rows if r['status'] == 'Unsettled']
    elif quick == 'partial':   rows = [r for r in all_rows if r['status'] == 'Partial']
    elif quick == 'settled':   rows = [r for r in all_rows if r['status'] == 'Settled']
    else:                      rows = all_rows

    # ── Batch-load checkout staff from AuditLog ───────────────────────────────
    res_ids = [r['reservation'].id for r in rows]
    checkout_user_map = {}
    if res_ids:
        audit_entries = (AuditLog.query
                         .filter(AuditLog.entity_type == 'Reservation',
                                 AuditLog.action == 'checkout',
                                 AuditLog.entity_id.in_(res_ids))
                         .with_entities(AuditLog.entity_id, AuditLog.staff_user_id)
                         .all())
        res_to_staff = {eid: sid for eid, sid in audit_entries}
        staff_ids    = set(res_to_staff.values())
        user_map     = {u.id: u.full_name
                        for u in User.query.filter(User.id.in_(staff_ids)).all()} if staff_ids else {}
        checkout_user_map = {eid: user_map.get(sid, '—')
                             for eid, sid in res_to_staff.items()}

    for r in rows:
        r['checkout_by'] = checkout_user_map.get(r['reservation'].id, '—')

    # ── Excel export ──────────────────────────────────────────────────────────
    if fmt == 'excel':
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response('openpyxl not installed.', status=500)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Invoice Register'

        hdr_fill = PatternFill('solid', fgColor='1F4E79')
        hdr_font = Font(color='FFFFFF', bold=True)

        label = f'Invoice Register — {from_date.strftime("%d %b %Y")} to {to_date.strftime("%d %b %Y")}'
        ws.append([label])
        ws.cell(1, 1).font = Font(bold=True, size=13)
        ws.append([f'Tab: {quick.capitalize()}  |  Rows: {len(rows)}'])
        ws.append([])

        headers = [
            'Invoice No.', 'Checkout Date', 'Checkout Time', 'Guest Name', 'Phone',
            'Room', 'Booking Type', 'Arrival', 'Departure', 'Nights',
            'Room Charges (₹)', 'Extra Charges (₹)', 'Bill Total (₹)',
            'Paid (₹)', 'Balance (₹)', 'Status', 'Payment Modes', 'Checkout By',
        ]
        hr = ws.max_row + 1
        ws.append(headers)
        for ci, _ in enumerate(headers, 1):
            c = ws.cell(hr, ci)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal='center')

        for r in rows:
            res   = r['reservation']
            b     = r['billing']
            nights = (res.departure_date - res.arrival_date).days
            co_date = res.checked_out_at.date().isoformat() if res.checked_out_at else ''
            co_time = res.checked_out_at.strftime('%H:%M')  if res.checked_out_at else ''
            modes_str = ', '.join(f'{k}:{v:,.0f}' for k, v in r['modes'].items())
            ws.append([
                res.booking_reference or f'RES-{res.id}',
                co_date, co_time,
                res.guest.name  if res.guest else '',
                res.guest.phone if res.guest else '',
                res.room.room_number if res.room else '',
                res.booking_type or '',
                res.arrival_date.isoformat() if res.arrival_date else '',
                res.departure_date.isoformat() if res.departure_date else '',
                nights,
                round(float(res.rate_per_night or 0) * nights, 2),
                round(b['total'] - float(res.rate_per_night or 0) * nights, 2),
                round(b['total'], 2),
                round(b['paid'], 2),
                round(b['balance'], 2),
                r['status'],
                modes_str,
                r['checkout_by'],
            ])

        # Auto-width
        for ci in range(1, len(headers) + 1):
            max_len = max(
                (len(str(ws.cell(row, ci).value or '')) for row in range(1, ws.max_row + 1)),
                default=10
            )
            ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 4, 40)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        tab_label = quick.capitalize()
        fname = f'invoice_register_{tab_label}_{from_date.isoformat()}_{to_date.isoformat()}.xlsx'
        from flask import send_file as _sf
        return _sf(buf, download_name=fname, as_attachment=True,
                   mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    return render_template(
        'billing/invoice_register.html',
        rows=rows,
        summary=summary,
        quick=quick,
        from_date=from_date,
        to_date=to_date,
        search=search,
        room_f=room_f,
        btype_f=btype_f,
    )


# ---------------------------------------------------------------------------
# Void Request — FrontDesk creates request
# ---------------------------------------------------------------------------

@billing_bp.route('/void/request/<int:payment_id>', methods=['POST'])
@login_required
@limiter.limit('20 per minute; 60 per hour')
def request_void(payment_id):
    if not payment_void_service.can_request_void(current_user):
        flash('You do not have permission to request a void.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    reason = request.form.get('reason', '').strip()
    if not reason:
        flash('Void reason is required.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    # Audit-lock guard: cannot request void on a payment that belongs to a
    # closed Night Audit date (Front Desk staff cannot override).
    payment = Payment.query.get_or_404(payment_id)
    from app.services import assert_business_date_unlocked
    ok, lock_err = assert_business_date_unlocked(payment.payment_date,
                                                  'request void')
    if not ok:
        flash(lock_err, 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    result = payment_void_service.request_void(
        payment_id=payment_id,
        user=current_user,
        reason=reason,
        ip_address=request.remote_addr,
    )
    flash(result.message, 'success' if result.success else 'danger')
    return redirect(request.referrer or url_for('main.reservations'))


# ---------------------------------------------------------------------------
# Void Dashboard — Manager/Admin sees all pending requests
# ---------------------------------------------------------------------------

@billing_bp.route('/void/pending')
@_manager_required
def void_pending():
    pending = payment_void_service.get_pending_void_requests()
    return render_template('billing/void_pending.html', pending=pending)


# ---------------------------------------------------------------------------
# Approve void request
# ---------------------------------------------------------------------------

@billing_bp.route('/void/approve/<int:void_request_id>', methods=['POST'])
@_manager_required
@limiter.limit('30 per minute')
def approve_void(void_request_id):
    result = payment_void_service.approve_void(
        void_request_id=void_request_id,
        approver=current_user,
        ip_address=request.remote_addr,
    )
    flash(result.message, 'success' if result.success else 'danger')
    return redirect(url_for('billing.void_pending'))


# ---------------------------------------------------------------------------
# Reject void request
# ---------------------------------------------------------------------------

@billing_bp.route('/void/reject/<int:void_request_id>', methods=['POST'])
@_manager_required
@limiter.limit('30 per minute')
def reject_void(void_request_id):
    rejection_reason = request.form.get('rejection_reason', '').strip()
    result = payment_void_service.reject_void(
        void_request_id=void_request_id,
        approver=current_user,
        rejection_reason=rejection_reason,
        ip_address=request.remote_addr,
    )
    flash(result.message, 'success' if result.success else 'danger')
    return redirect(url_for('billing.void_pending'))


# ---------------------------------------------------------------------------
# Direct void — Manager/Admin single step (replaces old /payment/<id>/void)
# ---------------------------------------------------------------------------

@billing_bp.route('/void/direct/<int:payment_id>', methods=['POST'])
@_manager_required
@limiter.limit('20 per minute; 60 per hour')
def direct_void(payment_id):
    reason = request.form.get('reason', '').strip()
    if not reason:
        flash('Void reason is mandatory.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    # Audit-lock guard with optional Admin override.
    payment = Payment.query.get_or_404(payment_id)
    if payment.is_correction:
        flash('Cannot void a correction entry directly. Post a fresh correction instead.', 'warning')
        return redirect(request.referrer or url_for('main.reservations'))
    from app.services import (assert_or_admin_override, post_payment_correction,
                              get_locking_audit)
    ok, lock_err, override_used, override_reason = assert_or_admin_override(
        payment.payment_date, 'void payment', request, current_user
    )
    if not ok:
        flash(lock_err, 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    # ── Correction-entry path ─────────────────────────────────────────
    # Locked-date + Admin override → post reversal (+ optional replacement)
    # so the original Payment row stays untouched and audit-safe.
    is_locked = get_locking_audit(payment.payment_date) is not None
    if is_locked and override_used:
        from app.routes import _write_audit
        try:
            new_amount_raw = (request.form.get('correction_new_amount') or '').strip()
            new_amount = float(new_amount_raw) if new_amount_raw else None
            if new_amount is not None and new_amount <= 0:
                new_amount = None
            new_mode_raw = (request.form.get('correction_new_mode_id') or '').strip()
            new_mode_id = int(new_mode_raw) if new_mode_raw else None
            new_ref = (request.form.get('correction_new_reference') or '').strip() or None
            full_reason = f'{reason} | Override: {override_reason}'
            corr = post_payment_correction(
                payment,
                new_amount=new_amount,
                new_mode_id=new_mode_id,
                new_reference=new_ref,
                reason=full_reason,
                user_id=current_user.id,
                audit_writer=_write_audit,
            )
            _write_audit('Payment', payment_id, 'audit_lock_override',
                         {'payment_date': payment.payment_date.isoformat(),
                          'amount': float(payment.amount)},
                         {'action':         'correction_pair',
                          'reversal_id':    corr['reversal'].id,
                          'replacement_id': (corr['replacement'].id
                                              if corr['replacement'] else None),
                          'override_reason': override_reason,
                          'admin_user_id':   current_user.id})
            db.session.commit()
        except Exception as exc:
            db.session.rollback()
            flash(f'Correction failed: {exc}', 'danger')
            return redirect(request.referrer or url_for('main.reservations'))
        msg = (f'Correction posted — reversal ₹{float(payment.amount):,.2f}'
               + (f' + replacement ₹{float(corr["replacement"].amount):,.2f}'
                  if corr['replacement'] else '')
               + '. Original row preserved.')
        flash(msg, 'success')
        return redirect(request.referrer or url_for('main.reservations'))

    # ── Standard void path (open audit date — direct mutation is fine) ──
    result = payment_void_service.direct_void(
        payment_id=payment_id,
        user=current_user,
        reason=reason,
        ip_address=request.remote_addr,
    )
    flash(result.message, 'success' if result.success else 'danger')
    return redirect(request.referrer or url_for('main.reservations'))


# ---------------------------------------------------------------------------
# AJAX: void eligibility check (for UI to show/hide void button)
# ---------------------------------------------------------------------------

@billing_bp.route('/void/eligible/<int:payment_id>')
@login_required
def void_eligible(payment_id):
    payment = db.session.get(Payment, payment_id)
    if not payment:
        return jsonify({'eligible': False, 'reason': 'Payment not found.'})
    ok, msg = payment_void_service.validate_void_eligibility(payment, current_user)
    return jsonify({'eligible': ok, 'reason': msg})


# ---------------------------------------------------------------------------
# Void history for a reservation
# ---------------------------------------------------------------------------

@billing_bp.route('/void/history/<int:reservation_id>')
@_manager_required
def void_history(reservation_id):
    from app.models import Reservation
    reservation = db.session.get(Reservation, reservation_id)
    if not reservation:
        flash('Reservation not found.', 'danger')
        return redirect(url_for('main.reservations'))

    requests = (
        VoidRequest.query
        .join(Payment)
        .filter(Payment.reservation_id == reservation_id)
        .order_by(VoidRequest.requested_at.desc())
        .all()
    )
    return render_template(
        'billing/void_history.html',
        reservation=reservation,
        void_requests=requests,
    )


# ---------------------------------------------------------------------------
# Credit Notes
# ---------------------------------------------------------------------------

@billing_bp.route('/credit-notes')
@login_required
def credit_notes_list():
    """List all credit notes."""
    if current_user.role not in ('Admin', 'Manager', 'Accountant'):
        flash('Access denied.', 'danger')
        return redirect(url_for('main.dashboard'))

    from app.models import CreditNote
    notes = (
        CreditNote.query
        .order_by(CreditNote.issued_at.desc())
        .all()
    )
    return render_template('billing/credit_notes.html', notes=notes)


@billing_bp.route('/credit-note/<int:reservation_id>', methods=['POST'])
@_manager_required
@limiter.limit('20 per minute')
def create_credit_note(reservation_id):
    """Create a credit note for a reservation (Manager/Admin only)."""
    from datetime import datetime
    from decimal import Decimal, ROUND_HALF_UP
    from app.models import CreditNote, Reservation, Settings

    reservation = db.session.get(Reservation, reservation_id)
    if not reservation:
        flash('Reservation not found.', 'danger')
        return redirect(url_for('main.reservations'))

    if not reservation.invoice_number:
        flash('No invoice found for this reservation. Cannot issue credit note.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    reason = request.form.get('reason', '').strip()
    if not reason:
        flash('Reason is required for a credit note.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    try:
        taxable = Decimal(request.form.get('taxable_amount', '0'))
        cgst    = Decimal(request.form.get('cgst_amount', '0'))
        sgst    = Decimal(request.form.get('sgst_amount', '0'))
        igst    = Decimal(request.form.get('igst_amount', '0'))
    except Exception:
        flash('Invalid amount values.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    TWO = Decimal('0.01')
    taxable = taxable.quantize(TWO, rounding=ROUND_HALF_UP)
    cgst    = cgst.quantize(TWO, rounding=ROUND_HALF_UP)
    sgst    = sgst.quantize(TWO, rounding=ROUND_HALF_UP)
    igst    = igst.quantize(TWO, rounding=ROUND_HALF_UP)
    total   = (taxable + cgst + sgst + igst).quantize(TWO)

    if total <= 0:
        flash('Credit note total must be greater than zero.', 'danger')
        return redirect(request.referrer or url_for('main.reservations'))

    notes_text = request.form.get('notes', '').strip()

    # Generate sequential credit note number: CN-YYYY-000001
    max_retries = 3
    for attempt in range(max_retries):
        try:
            counter_row = (
                db.session.query(Settings)
                .filter_by(key='credit_note_counter')
                .with_for_update()
                .first()
            )
            if counter_row:
                new_val = int(counter_row.value or 0) + 1
                db.session.execute(
                    db.update(Settings)
                    .where(Settings.key == 'credit_note_counter')
                    .where(Settings.value == counter_row.value)
                    .values(value=str(new_val))
                )
                current_n = new_val
            else:
                current_n = 1
                db.session.add(Settings(
                    key='credit_note_counter', value='1',
                    description='Running sequential credit note counter'))
                db.session.flush()

            year = datetime.utcnow().strftime('%Y')
            cn_number = f'CN-{year}-{current_n:06d}'

            cn = CreditNote(
                credit_note_number=cn_number,
                reservation_id=reservation.id,
                original_invoice_number=reservation.invoice_number,
                reason=reason,
                taxable_amount=taxable,
                cgst_amount=cgst,
                sgst_amount=sgst,
                igst_amount=igst,
                total_amount=total,
                issued_by_user_id=current_user.id,
                notes=notes_text or None,
            )
            db.session.add(cn)
            db.session.commit()
            flash(f'Credit note {cn_number} issued successfully.', 'success')
            return redirect(url_for('billing.credit_notes_list'))

        except Exception as e:
            db.session.rollback()
            if attempt == max_retries - 1:
                flash(f'Failed to create credit note: {e}', 'danger')
                return redirect(request.referrer or url_for('main.reservations'))

    return redirect(request.referrer or url_for('main.reservations'))
